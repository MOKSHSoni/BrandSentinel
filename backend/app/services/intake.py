"""Content intake: plain text, DOCX and PDF become one normalised text string."""

import io
import re

MAX_BYTES = 5 * 1024 * 1024
MAX_CHARS = 20_000


class IntakeError(ValueError):
    pass


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n\n".join(p for p in parts if p.strip())


def _pdf_text(data: bytes) -> str:
    import fitz

    with fitz.open(stream=data, filetype="pdf") as pdf:
        return "\n\n".join(page.get_text() for page in pdf)


def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_BYTES:
        raise IntakeError("File is larger than 5 MB.")
    name = filename.lower()
    try:
        if name.endswith(".docx"):
            text = _docx_text(data)
        elif name.endswith(".pdf"):
            text = _pdf_text(data)
        elif name.endswith((".txt", ".md")):
            text = data.decode("utf-8", errors="replace")
        else:
            raise IntakeError("Upload a .docx, .pdf or .txt file.")
    except IntakeError:
        raise
    except Exception as e:
        raise IntakeError(f"Could not read {filename}: {e}") from e
    return normalise(text)


def normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise IntakeError("The draft is empty.")
    if len(text) > MAX_CHARS:
        raise IntakeError(f"Drafts are limited to {MAX_CHARS:,} characters.")
    return text


def make_docx(title: str, content: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.core_properties.title = title
    for para in content.split("\n\n"):
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
